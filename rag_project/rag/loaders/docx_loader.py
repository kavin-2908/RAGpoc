"""
DOCX Loader — Extracts text from Word documents using python-docx.

python-docx reads .docx files (the modern Word format, not the old .doc).
It provides access to paragraphs, tables, headers, and footers.

For V1, we extract paragraph text only. Tables and headers can be
added in a future version if needed.
"""

import logging

from docx import Document

logger = logging.getLogger(__name__)


class DOCXLoader:
    """Extracts text from DOCX (Word) files."""

    def load(self, file_path: str) -> str:
        """
        Extract text from a DOCX file.

        Reads all paragraphs from the document and joins non-empty
        ones with newlines.

        Args:
            file_path: Absolute path to the DOCX file.

        Returns:
            Extracted text as a single string.

        Raises:
            Exception: If the DOCX file is corrupted or cannot be opened.
        """
        logger.info("Loading DOCX: %s", file_path)

        # Document() opens and parses the .docx file
        doc = Document(file_path)

        # doc.paragraphs is a list of Paragraph objects
        # Each paragraph has a .text attribute with its content
        text_parts: list[str] = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        full_text = "\n".join(text_parts)
        logger.info(
            "DOCX loaded: %d paragraphs, %d characters total",
            len(text_parts),
            len(full_text),
        )

        return full_text
