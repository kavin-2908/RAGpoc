"""
Excel Loader — Extracts text from .xlsx files using openpyxl.

openpyxl reads the modern Excel format (.xlsx). It does NOT support
the legacy .xls format.

Strategy:
- Reads every sheet in the workbook.
- For each sheet, reads every row.
- Each row is joined into a comma-separated string.
- Sheet names are included as section headers for context.

This produces a readable text representation where the LLM
can understand which data came from which sheet.
"""

import logging

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


class ExcelLoader:
    """Extracts text from Excel (.xlsx) files."""

    def load(self, file_path: str) -> str:
        """
        Read an Excel file and convert it to a text representation.

        Iterates over all sheets. For each sheet, reads all rows
        and joins cell values with commas. Sheets are separated
        by a header line with the sheet name.

        Args:
            file_path: Absolute path to the .xlsx file.

        Returns:
            Text representation of the Excel data.

        Raises:
            Exception: If the file is corrupted or cannot be opened.
        """
        logger.info("Loading Excel: %s", file_path)

        # data_only=True reads computed values instead of formulas
        workbook = load_workbook(file_path, data_only=True)

        text_parts: list[str] = []

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]

            # Add sheet name as a header for context
            text_parts.append(f"--- Sheet: {sheet_name} ---")

            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                # Convert each cell to string, skip None values
                cell_values = [
                    str(cell).strip()
                    for cell in row
                    if cell is not None
                ]

                if cell_values:
                    text_parts.append(", ".join(cell_values))
                    row_count += 1

            logger.debug(
                "Sheet '%s': %d rows extracted", sheet_name, row_count
            )

        workbook.close()

        full_text = "\n".join(text_parts)
        logger.info(
            "Excel loaded: %d sheets, %d characters total",
            len(workbook.sheetnames),
            len(full_text),
        )

        return full_text
